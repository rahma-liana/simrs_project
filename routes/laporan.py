import io
import csv
from datetime import datetime, timedelta
from flask import Blueprint, render_template, request, Response, send_file
from database import query_db
from helpers import login_required, get_hospital_profile, format_rupiah, format_tanggal_indo
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

laporan_bp = Blueprint('laporan', __name__, url_prefix='/laporan')

@laporan_bp.route('/')
@login_required
def index():
    return render_template('laporan/index.html')

# -------------------------------------------------------------
# 1. LAPORAN PASIEN
# -------------------------------------------------------------
@laporan_bp.route('/pasien')
@login_required
def pasien():
    gol_darah = request.args.get('gol_darah', '')
    jenis_kelamin = request.args.get('jenis_kelamin', '')
    export = request.args.get('export', '')
    
    query = "SELECT * FROM pasien WHERE 1=1"
    params = []
    if gol_darah:
        query += " AND golongan_darah = %s"
        params.append(gol_darah)
    if jenis_kelamin:
        query += " AND jenis_kelamin = %s"
        params.append(jenis_kelamin)
    query += " ORDER BY id DESC"
    
    data = query_db(query, tuple(params))
    hosp = get_hospital_profile()
    
    if export == 'excel':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Laporan Pasien"
        
        headers = ["No", "No. RM", "NIK", "Nama Lengkap", "L/P", "Tempat Lahir", "Tgl Lahir", "Gol. Darah", "No. Telepon", "Alamat"]
        ws.append([hosp['nama_rs']])
        ws.append(["LAPORAN DATA PASIEN"])
        ws.append([f"Dicetak pada: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.append([])
        ws.append(headers)
        
        # Style header
        header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=5, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            
        for idx, row in enumerate(data, 1):
            ws.append([
                idx, row['no_rm'], row['nik'], row['nama_lengkap'], row['jenis_kelamin'],
                row['tempat_lahir'], str(row['tanggal_lahir']), row['golongan_darah'],
                row['no_telepon'], row['alamat']
            ])
            
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, download_name=f"Laporan_Pasien_{datetime.now().strftime('%Y%m%d')}.xlsx", as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        
    if export == 'print':
        return render_template('laporan/cetak_pasien.html', data=data, hosp=hosp, filters={'gol_darah': gol_darah, 'jenis_kelamin': jenis_kelamin})
        
    return render_template('laporan/pasien.html', data=data, gol_darah=gol_darah, jenis_kelamin=jenis_kelamin)

# -------------------------------------------------------------
# 2. LAPORAN KUNJUNGAN PASIEN
# -------------------------------------------------------------
@laporan_bp.route('/kunjungan')
@login_required
def kunjungan():
    tgl_mulai = request.args.get('tgl_mulai', (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d'))
    tgl_selesai = request.args.get('tgl_selesai', datetime.now().strftime('%Y-%m-%d'))
    poli_id = request.args.get('poli_id', '')
    status = request.args.get('status', '')
    export = request.args.get('export', '')
    
    query = """
        SELECT p.*, pas.nama_lengkap as nama_pasien, pas.no_rm, 
               d.nama_dokter, pol.nama_poli
        FROM pendaftaran p
        JOIN pasien pas ON p.pasien_id = pas.id
        JOIN dokter d ON p.dokter_id = d.id
        JOIN poli pol ON p.poli_id = pol.id
        WHERE p.tanggal_kunjungan BETWEEN %s AND %s
    """
    params = [tgl_mulai, tgl_selesai]
    if poli_id:
        query += " AND p.poli_id = %s"
        params.append(poli_id)
    if status:
        query += " AND p.status_antrean = %s"
        params.append(status)
    query += " ORDER BY p.tanggal_kunjungan DESC, p.no_antrean ASC"
    
    data = query_db(query, tuple(params))
    poli_list = query_db("SELECT * FROM poli ORDER BY nama_poli ASC")
    hosp = get_hospital_profile()
    
    if export == 'excel':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Laporan Kunjungan"
        
        headers = ["No", "Tgl Kunjungan", "No. Reg", "No. Antrean", "No. RM", "Nama Pasien", "Poli", "Dokter", "Status", "Keluhan"]
        ws.append([hosp['nama_rs']])
        ws.append([f"LAPORAN KUNJUNGAN PASIEN ({tgl_mulai} s/d {tgl_selesai})"])
        ws.append([])
        ws.append(headers)
        
        header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            
        for idx, row in enumerate(data, 1):
            ws.append([
                idx, str(row['tanggal_kunjungan']), row['no_pendaftaran'], row['no_antrean'],
                row['no_rm'], row['nama_pasien'], row['nama_poli'], row['nama_dokter'],
                row['status_antrean'], row['keluhan_awal'] or '-'
            ])
            
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, download_name=f"Laporan_Kunjungan_{datetime.now().strftime('%Y%m%d')}.xlsx", as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    if export == 'print':
        return render_template('laporan/cetak_kunjungan.html', data=data, hosp=hosp, tgl_mulai=tgl_mulai, tgl_selesai=tgl_selesai)

    return render_template('laporan/kunjungan.html', data=data, poli_list=poli_list, 
                           tgl_mulai=tgl_mulai, tgl_selesai=tgl_selesai, selected_poli=poli_id, selected_status=status)

# -------------------------------------------------------------
# 3. LAPORAN STOK & MUTASI OBAT (FARMASI)
# -------------------------------------------------------------
@laporan_bp.route('/farmasi')
@login_required
def farmasi():
    tipe = request.args.get('tipe', 'stok') # 'stok' atau 'mutasi'
    export = request.args.get('export', '')
    hosp = get_hospital_profile()
    
    if tipe == 'mutasi':
        tgl_mulai = request.args.get('tgl_mulai', (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d'))
        tgl_selesai = request.args.get('tgl_selesai', datetime.now().strftime('%Y-%m-%d'))
        jenis = request.args.get('jenis', '')
        
        query = """
            SELECT m.*, o.kode_obat, o.nama_obat, o.satuan, o.kategori
            FROM obat_mutasi m
            JOIN obat o ON m.obat_id = o.id
            WHERE m.tanggal BETWEEN %s AND %s
        """
        params = [tgl_mulai, tgl_selesai]
        if jenis:
            query += " AND m.jenis_mutasi = %s"
            params.append(jenis)
        query += " ORDER BY m.tanggal DESC, m.id DESC"
        
        data = query_db(query, tuple(params))
        
        if export == 'excel':
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Mutasi Obat"
            ws.append([hosp['nama_rs']])
            ws.append([f"LAPORAN MUTASI OBAT ({tgl_mulai} s/d {tgl_selesai})"])
            ws.append([])
            headers = ["No", "Tanggal", "Kode Obat", "Nama Obat", "Kategori", "Jenis", "Jumlah", "Stok Sebelum", "Stok Sesudah", "Referensi", "Keterangan"]
            ws.append(headers)
            for idx, r in enumerate(data, 1):
                ws.append([idx, str(r['tanggal']), r['kode_obat'], r['nama_obat'], r['kategori'], r['jenis_mutasi'], r['jumlah'], r['stok_sebelum'], r['stok_sesudah'], r['no_referensi'], r['keterangan']])
            out = io.BytesIO()
            wb.save(out)
            out.seek(0)
            return send_file(out, download_name=f"Laporan_Mutasi_Obat_{datetime.now().strftime('%Y%m%d')}.xlsx", as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            
        if export == 'print':
            return render_template('laporan/cetak_mutasi_obat.html', data=data, hosp=hosp, tgl_mulai=tgl_mulai, tgl_selesai=tgl_selesai)
            
        return render_template('laporan/farmasi_mutasi.html', data=data, tgl_mulai=tgl_mulai, tgl_selesai=tgl_selesai, jenis=jenis)
        
    else: # stok
        kategori = request.args.get('kategori', '')
        query = "SELECT * FROM obat WHERE 1=1"
        params = []
        if kategori:
            query += " AND kategori = %s"
            params.append(kategori)
        query += " ORDER BY kategori, nama_obat ASC"
        data = query_db(query, tuple(params))
        kategori_list = query_db("SELECT DISTINCT kategori FROM obat ORDER BY kategori ASC")
        
        if export == 'excel':
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Stok Obat"
            ws.append([hosp['nama_rs']])
            ws.append(["LAPORAN STOK PERSEDIAAN OBAT & FARMASI"])
            ws.append([])
            headers = ["No", "Kode Obat", "Nama Obat", "Kategori", "Satuan", "Harga Beli", "Harga Jual", "Stok Saat Ini", "Stok Min", "Kadaluwarsa", "Status"]
            ws.append(headers)
            for idx, r in enumerate(data, 1):
                stt = "MENIPIS" if r['stok_saat_ini'] <= r['stok_minimum'] else "Aman"
                ws.append([idx, r['kode_obat'], r['nama_obat'], r['kategori'], r['satuan'], float(r['harga_beli']), float(r['harga_jual']), r['stok_saat_ini'], r['stok_minimum'], str(r['tanggal_kadaluwarsa']), stt])
            out = io.BytesIO()
            wb.save(out)
            out.seek(0)
            return send_file(out, download_name=f"Laporan_Stok_Obat_{datetime.now().strftime('%Y%m%d')}.xlsx", as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        if export == 'print':
            return render_template('laporan/cetak_stok_obat.html', data=data, hosp=hosp)

        return render_template('laporan/farmasi_stok.html', data=data, kategori_list=kategori_list, selected_kategori=kategori)

# -------------------------------------------------------------
# 4. LAPORAN TRANSAKSI & PENDAPATAN (KEUANGAN)
# -------------------------------------------------------------
@laporan_bp.route('/pendapatan')
@login_required
def pendapatan():
    tgl_mulai = request.args.get('tgl_mulai', (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d'))
    tgl_selesai = request.args.get('tgl_selesai', datetime.now().strftime('%Y-%m-%d'))
    metode = request.args.get('metode', '')
    export = request.args.get('export', '')
    
    query = """
        SELECT pem.*, t.no_tagihan, t.biaya_pendaftaran, t.biaya_konsultasi, 
               t.biaya_tindakan, t.biaya_obat, t.total_tagihan,
               pas.nama_lengkap as nama_pasien, pas.no_rm, pol.nama_poli
        FROM pembayaran pem
        JOIN tagihan t ON pem.tagihan_id = t.id
        JOIN pendaftaran p ON t.pendaftaran_id = p.id
        JOIN pasien pas ON t.pasien_id = pas.id
        JOIN poli pol ON p.poli_id = pol.id
        WHERE DATE(pem.tanggal_bayar) BETWEEN %s AND %s
    """
    params = [tgl_mulai, tgl_selesai]
    if metode:
        query += " AND pem.metode_pembayaran = %s"
        params.append(metode)
    query += " ORDER BY pem.tanggal_bayar DESC"
    
    data = query_db(query, tuple(params))
    total_pendapatan = sum(float(r['jumlah_bayar']) for r in data)
    hosp = get_hospital_profile()
    
    if export == 'excel':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Laporan Pendapatan"
        ws.append([hosp['nama_rs']])
        ws.append([f"LAPORAN TRANSAKSI & PENDAPATAN RUMAH SAKIT ({tgl_mulai} s/d {tgl_selesai})"])
        ws.append([])
        headers = ["No", "Waktu Transaksi", "No. Kuitansi", "No. Tagihan", "No. RM", "Nama Pasien", "Poli", "Metode", "Jumlah Dibayar (Rp)"]
        ws.append(headers)
        
        for idx, r in enumerate(data, 1):
            ws.append([
                idx, str(r['tanggal_bayar']), r['no_pembayaran'], r['no_tagihan'],
                r['no_rm'], r['nama_pasien'], r['nama_poli'], r['metode_pembayaran'], float(r['jumlah_bayar'])
            ])
        ws.append([])
        ws.append(["TOTAL PENDAPATAN", "", "", "", "", "", "", "", total_pendapatan])
        
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, download_name=f"Laporan_Pendapatan_{datetime.now().strftime('%Y%m%d')}.xlsx", as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    if export == 'print':
        return render_template('laporan/cetak_pendapatan.html', data=data, hosp=hosp, total_pendapatan=total_pendapatan,
                               tgl_mulai=tgl_mulai, tgl_selesai=tgl_selesai)

    return render_template('laporan/pendapatan.html', data=data, total_pendapatan=total_pendapatan,
                           tgl_mulai=tgl_mulai, tgl_selesai=tgl_selesai, selected_metode=metode)
